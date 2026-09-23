"""Condense the corpus folder into one candidate profile, cached until the corpus changes."""
from __future__ import annotations

import hashlib
from pathlib import Path

from .config import ROOT
from .llm import Gemini

PROFILE_PATH = ROOT / "profile.md"
HASH_PATH = ROOT / ".cache" / "profile.hash"
EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}
# About 100k tokens: well under the free tier's per-minute token limit.
MAX_CORPUS_CHARS = 400_000

SYSTEM = """You turn a candidate's career documents into a factual profile that a job-matching \
screener will use to judge fit. Use only facts stated in the documents; never infer or embellish. \
Where the documents disagree, prefer the most recent one."""

PROMPT = """Write the candidate profile in Markdown with these sections:

## Summary
Two or three sentences: current title, total years of experience, core specialty.
## Current role
Title, company, start date, responsibilities, and scope (team size, systems, budget where stated).
## Work history
For each role, newest first: title, company, dates, and 2-5 achievements, keeping every metric.
## Skills
Group into languages, frameworks/libraries, cloud/infrastructure/tools, methods, and domains. \
After each skill, add in brackets where it was used and for roughly how long, e.g. "Python (5 yrs; Acme, thesis)".
## Projects
Name, one-line purpose, stack, and outcome.
## Education and certifications
## Work authorisation, location and preferences
Only if stated.

Keep it under about 1,500 words. Documents follow.

{docs}"""


def _read(path: Path) -> str:
    if path.suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if path.suffix == ".docx":
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in table.rows]
        return "\n".join(parts)
    return path.read_text(encoding="utf-8", errors="replace")


def read_corpus(corpus_dir: Path) -> list[tuple[str, str]]:
    if not corpus_dir.is_dir():
        raise SystemExit(f"Corpus folder {corpus_dir} does not exist.")
    docs = []
    for path in sorted(corpus_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in EXTENSIONS and path.name != "README.md":
            text = _read(path).strip()
            if text:
                docs.append((str(path.relative_to(corpus_dir)), text))
    if not docs:
        raise SystemExit(f"No .pdf/.docx/.md/.txt documents found in {corpus_dir}.")
    return docs


def load_profile(corpus_dir: Path, llm: Gemini, model: str, force: bool = False) -> str:
    docs = read_corpus(corpus_dir)
    digest = hashlib.sha256("\0".join(f"{n}\0{t}" for n, t in docs).encode()).hexdigest()
    if not force and PROFILE_PATH.exists() and HASH_PATH.exists() and HASH_PATH.read_text() == digest:
        return PROFILE_PATH.read_text(encoding="utf-8")

    joined = "\n\n".join(f"=== {name} ===\n{text}" for name, text in docs)
    if len(joined) > MAX_CORPUS_CHARS:
        raise SystemExit(f"Corpus is {len(joined):,} characters, over the {MAX_CORPUS_CHARS:,} that fit "
                         f"in one free-tier request. Remove or trim documents in {corpus_dir}.")
    print(f"Building candidate profile from {len(docs)} document(s) with {model}...")
    profile = llm.generate(model, SYSTEM, PROMPT.format(docs=joined)).strip()
    PROFILE_PATH.write_text(profile + "\n", encoding="utf-8")
    HASH_PATH.parent.mkdir(exist_ok=True)
    HASH_PATH.write_text(digest)
    print(f"Wrote {PROFILE_PATH}. Review it; the matcher treats it as the truth about you.")
    return profile
